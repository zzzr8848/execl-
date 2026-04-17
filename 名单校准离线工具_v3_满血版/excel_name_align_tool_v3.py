
import tkinter as tk
from tkinter import ttk, filedialog, messagebox
from pathlib import Path
import traceback
import unicodedata
import json
import re
from collections import defaultdict, Counter
from difflib import SequenceMatcher
from datetime import datetime

try:
    from openpyxl import load_workbook, Workbook
    from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
except Exception:
    raise SystemExit("缺少依赖 openpyxl。请先安装：pip install openpyxl")


APP_TITLE = "Excel名单校准离线工具 v3 满血版"
NO_SECONDARY = "（不使用）"

MATCH_MODE_PRIMARY = "仅主键精确匹配"
MATCH_MODE_COMBINED_STRICT = "主键+辅键联合精确匹配"
MATCH_MODE_COMBINED_FALLBACK = "优先联合匹配，必要时回退主键"

STATUS_MATCH_PRIMARY = "已匹配（主键）"
STATUS_MATCH_COMBINED = "已匹配（联合键）"
STATUS_MATCH_FALLBACK = "已匹配（主键回退）"
STATUS_SRC_PRIMARY_EMPTY = "源名单主键为空"
STATUS_SRC_SECONDARY_EMPTY = "源名单辅键为空，无法联合匹配"
STATUS_PRIMARY_DUP = "主键重复，需人工核对"
STATUS_COMBINED_DUP = "联合键重复，需人工核对"
STATUS_MISSING = "缺漏（源名单有，新名单无）"

HEADER_FILL = PatternFill("solid", fgColor="D9E2F3")
GREEN_FILL = PatternFill("solid", fgColor="E2F0D9")
YELLOW_FILL = PatternFill("solid", fgColor="FFF2CC")
RED_FILL = PatternFill("solid", fgColor="FCE4D6")
GRAY_FILL = PatternFill("solid", fgColor="EDEDED")
BLUE_FILL = PatternFill("solid", fgColor="DDEBF7")

THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

NAME_KEYWORDS = ("姓名", "名字", "name")
IDENTITY_KEYWORDS = (
    "电话", "手机", "手机号", "联系电话",
    "身份证", "证件", "证号", "编号", "id", "ID"
)


def remove_invisible_and_spaces(s: str) -> str:
    s = unicodedata.normalize("NFKC", s)
    cleaned = []
    invisible_chars = {
        "\ufeff",  # BOM / ZWNBSP
        "\u200b",  # zero-width space
        "\u200c",
        "\u200d",
        "\u2060",
        "\u180e",
    }
    for ch in s:
        cat = unicodedata.category(ch)
        if ch in invisible_chars:
            continue
        if cat.startswith("Z") or cat == "Cf":
            continue
        if ch in ("\n", "\r", "\t"):
            continue
        cleaned.append(ch)
    return "".join(cleaned).strip()


def normalize_name(value) -> str:
    if value is None:
        return ""
    return remove_invisible_and_spaces(str(value))


def normalize_identifier(value) -> str:
    if value is None:
        return ""
    s = remove_invisible_and_spaces(str(value)).upper()
    return "".join(ch for ch in s if ch.isalnum())


def normalize_generic(value) -> str:
    if value is None:
        return ""
    return remove_invisible_and_spaces(str(value))


def normalize_by_header(value, header_name: str) -> str:
    header_text = "" if header_name is None else str(header_name)
    h = header_text.lower()
    if any(k.lower() in h for k in NAME_KEYWORDS):
        return normalize_name(value)
    if any(k.lower() in h for k in IDENTITY_KEYWORDS):
        return normalize_identifier(value)
    return normalize_generic(value)


def get_sheet_names(file_path):
    wb = load_workbook(file_path, read_only=True, data_only=True)
    names = wb.sheetnames
    wb.close()
    return names


def read_sheet_headers(file_path, sheet_name):
    wb = load_workbook(file_path, read_only=True, data_only=True)
    ws = wb[sheet_name]
    rows = ws.iter_rows(min_row=1, max_row=1, values_only=True)
    first = next(rows, None)
    wb.close()
    if not first:
        return []
    return ["" if v is None else str(v).strip() for v in first]


def read_rows(file_path, sheet_name):
    wb = load_workbook(file_path, read_only=True, data_only=True)
    ws = wb[sheet_name]
    all_rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if not all_rows:
        return [], []
    headers = ["" if v is None else str(v).strip() for v in all_rows[0]]
    data = []
    for row_idx, row in enumerate(all_rows[1:], start=2):
        item = {}
        for i, h in enumerate(headers):
            item[h] = row[i] if i < len(row) else None
        data.append({"_rownum": row_idx, "_data": item})
    return headers, data


def validate_headers(headers, label):
    if not headers:
        raise ValueError(f"{label}工作表为空，无法读取表头。")
    cleaned = [h for h in headers if h]
    if not cleaned:
        raise ValueError(f"{label}表头为空。")
    dup = [k for k, c in Counter(headers).items() if k and c > 1]
    if dup:
        raise ValueError(f"{label}存在重复表头：{', '.join(dup)}")


def build_row_metas(rows, primary_col, secondary_col):
    metas = []
    for idx, row in enumerate(rows, start=1):
        data = row["_data"]
        meta = {
            "id": idx,
            "rownum": row["_rownum"],
            "data": data,
            "primary_raw": data.get(primary_col),
            "secondary_raw": "" if secondary_col == NO_SECONDARY else data.get(secondary_col),
            "primary_norm": normalize_by_header(data.get(primary_col), primary_col),
            "secondary_norm": "" if secondary_col == NO_SECONDARY else normalize_by_header(data.get(secondary_col), secondary_col),
        }
        metas.append(meta)
    return metas


def make_index(metas, key_func):
    index = defaultdict(list)
    for meta in metas:
        key = key_func(meta)
        if key:
            index[key].append(meta["id"])
    return index


def combined_key(meta):
    if not meta["primary_norm"] or not meta["secondary_norm"]:
        return None
    return (meta["primary_norm"], meta["secondary_norm"])


def primary_key(meta):
    if not meta["primary_norm"]:
        return None
    return meta["primary_norm"]


def only_available(ids, used_new_ids):
    return [i for i in ids if i not in used_new_ids]


def last4(s: str) -> str:
    if not s:
        return ""
    digits = "".join(ch for ch in s if ch.isalnum())
    return digits[-4:] if len(digits) >= 4 else digits


def choose_match_for_source(
    src_meta,
    used_new_ids,
    source_primary_idx,
    new_primary_idx,
    source_combined_idx,
    new_combined_idx,
    mode,
):
    p = src_meta["primary_norm"]
    s = src_meta["secondary_norm"]

    if not p:
        return None, STATUS_SRC_PRIMARY_EMPTY

    if mode == MATCH_MODE_PRIMARY:
        src_p_ids = source_primary_idx.get(p, [])
        new_p_ids = new_primary_idx.get(p, [])
        if len(src_p_ids) > 1 or len(new_p_ids) > 1:
            return None, STATUS_PRIMARY_DUP
        candidates = only_available(new_p_ids, used_new_ids)
        if len(candidates) == 1:
            return candidates[0], STATUS_MATCH_PRIMARY
        return None, STATUS_MISSING

    if mode == MATCH_MODE_COMBINED_STRICT:
        if not s:
            return None, STATUS_SRC_SECONDARY_EMPTY
        ck = (p, s)
        src_c_ids = source_combined_idx.get(ck, [])
        new_c_ids = new_combined_idx.get(ck, [])
        if len(src_c_ids) > 1 or len(new_c_ids) > 1:
            return None, STATUS_COMBINED_DUP
        candidates = only_available(new_c_ids, used_new_ids)
        if len(candidates) == 1:
            return candidates[0], STATUS_MATCH_COMBINED
        return None, STATUS_MISSING

    # MATCH_MODE_COMBINED_FALLBACK
    if s:
        ck = (p, s)
        src_c_ids = source_combined_idx.get(ck, [])
        new_c_ids = new_combined_idx.get(ck, [])
        if len(src_c_ids) == 1 and len(new_c_ids) == 1:
            candidates = only_available(new_c_ids, used_new_ids)
            if len(candidates) == 1:
                return candidates[0], STATUS_MATCH_COMBINED
        elif len(src_c_ids) > 1 or len(new_c_ids) > 1:
            return None, STATUS_COMBINED_DUP

    src_p_ids = source_primary_idx.get(p, [])
    new_p_ids = new_primary_idx.get(p, [])
    if len(src_p_ids) > 1 or len(new_p_ids) > 1:
        return None, STATUS_PRIMARY_DUP
    candidates = only_available(new_p_ids, used_new_ids)
    if len(candidates) == 1:
        return candidates[0], STATUS_MATCH_FALLBACK if s else STATUS_MATCH_PRIMARY
    return None, STATUS_MISSING


def common_suffix_len(a: str, b: str) -> int:
    n = 0
    for x, y in zip(a[::-1], b[::-1]):
        if x == y:
            n += 1
        else:
            break
    return n


def score_possible_match(src_meta, new_meta):
    src_name = src_meta["primary_norm"]
    new_name = new_meta["primary_norm"]
    if not src_name or not new_name:
        return 0.0

    base = SequenceMatcher(None, src_name, new_name).ratio()
    suffix_len = common_suffix_len(src_name, new_name)
    same_first = src_name[:1] == new_name[:1]
    same_len = len(src_name) == len(new_name)

    aux_bonus = 0.0
    src_aux = src_meta["secondary_norm"]
    new_aux = new_meta["secondary_norm"]
    if src_aux and new_aux:
        if src_aux == new_aux:
            aux_bonus += 0.25
        elif last4(src_aux) and last4(src_aux) == last4(new_aux):
            aux_bonus += 0.15
        elif src_aux in new_aux or new_aux in src_aux:
            aux_bonus += 0.10

    plausible = (
        base >= 0.85
        or (same_first and base >= 0.65)
        or (same_len and suffix_len >= 2 and base >= 0.60)
        or (aux_bonus >= 0.15 and base >= 0.50)
    )
    if not plausible:
        return 0.0
    return min(1.0, base + aux_bonus)


def build_suspect_candidates(unmatched_source_metas, unmatched_new_metas, threshold=0.66, topn=3):
    if not unmatched_source_metas or not unmatched_new_metas:
        return []

    bucket_first = defaultdict(list)
    bucket_suffix2 = defaultdict(list)
    bucket_len = defaultdict(list)

    for nm in unmatched_new_metas:
        name = nm["primary_norm"]
        if not name:
            continue
        bucket_first[name[:1]].append(nm)
        bucket_suffix2[name[-2:] if len(name) >= 2 else name].append(nm)
        bucket_len[len(name)].append(nm)

    results = []
    for sm in unmatched_source_metas:
        src_name = sm["primary_norm"]
        if not src_name:
            continue

        candidates_pool = {}
        for nm in bucket_first.get(src_name[:1], []):
            candidates_pool[nm["id"]] = nm
        for nm in bucket_suffix2.get(src_name[-2:] if len(src_name) >= 2 else src_name, []):
            candidates_pool[nm["id"]] = nm
        for nm in bucket_len.get(len(src_name), []):
            candidates_pool[nm["id"]] = nm

        scored = []
        for nm in candidates_pool.values():
            score = score_possible_match(sm, nm)
            if score >= threshold:
                scored.append((score, nm))

        scored.sort(key=lambda x: (-x[0], x[1]["rownum"]))
        for rank, (score, nm) in enumerate(scored[:topn], start=1):
            results.append({
                "source_meta": sm,
                "new_meta": nm,
                "rank": rank,
                "score": round(score, 4),
            })
    return results


def autofit_columns(ws, max_width=40):
    for col in ws.columns:
        max_len = 0
        col_letter = col[0].column_letter
        for cell in col:
            try:
                v = "" if cell.value is None else str(cell.value)
            except Exception:
                v = ""
            max_len = max(max_len, len(v))
        ws.column_dimensions[col_letter].width = min(max(max_len + 2, 10), max_width)


def style_header(ws):
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BORDER


def apply_common_sheet_style(ws):
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for row in ws.iter_rows():
        for cell in row:
            cell.border = BORDER
            cell.alignment = Alignment(vertical="center")


def apply_status_fills(ws, status_col=1):
    for row in range(2, ws.max_row + 1):
        cell = ws.cell(row=row, column=status_col)
        text = "" if cell.value is None else str(cell.value)
        if text.startswith("已匹配"):
            cell.fill = GREEN_FILL
        elif "重复" in text:
            cell.fill = YELLOW_FILL
        elif "缺漏" in text or "为空" in text:
            cell.fill = RED_FILL
        else:
            cell.fill = GRAY_FILL


def style_summary_sheet(ws):
    for row in ws.iter_rows():
        for cell in row:
            cell.border = BORDER
            cell.alignment = Alignment(vertical="center")
    for cell in ws[1]:
        cell.fill = BLUE_FILL
        cell.font = Font(bold=True)
    ws.freeze_panes = "A2"
    autofit_columns(ws, max_width=55)


def build_summary_rows(summary):
    rows = [
        ["项目", "值"],
        ["生成时间", summary["generated_at"]],
        ["源名单文件", summary["source_file"]],
        ["新名单文件", summary["new_file"]],
        ["匹配模式", summary["match_mode"]],
        ["源名单主键列", summary["source_primary"]],
        ["新名单主键列", summary["new_primary"]],
        ["源名单辅键列", summary["source_secondary"]],
        ["新名单辅键列", summary["new_secondary"]],
        ["疑似匹配阈值", summary["threshold"]],
        ["源名单总人数", summary["source_total"]],
        ["新名单总人数", summary["new_total"]],
        ["已匹配（联合键）", summary["count_combined"]],
        ["已匹配（主键）", summary["count_primary"]],
        ["已匹配（主键回退）", summary["count_fallback"]],
        ["缺漏", summary["count_missing"]],
        ["主键重复需人工核对", summary["count_primary_dup"]],
        ["联合键重复需人工核对", summary["count_combined_dup"]],
        ["源名单主键为空", summary["count_primary_empty"]],
        ["源名单辅键为空", summary["count_secondary_empty"]],
        ["新名单多出（未被匹配）", summary["new_unmatched"]],
        ["疑似匹配候选条数", summary["suspect_count"]],
        ["提醒", "若存在重名，建议改用“姓名+电话后4位/身份证后4位”等联合键。"],
    ]
    for warning in summary.get("warnings", []):
        rows.append(["预检查提醒", warning])
    return rows


def safe_save_workbook(out_wb, output_file):
    target = Path(output_file)
    target.parent.mkdir(parents=True, exist_ok=True)

    try:
        out_wb.save(target)
        return str(target)
    except PermissionError:
        pass

    stem = target.stem
    suffix = target.suffix or ".xlsx"
    for i in range(1, 100):
        candidate = target.with_name(f"{stem}_{i}{suffix}")
        try:
            out_wb.save(candidate)
            return str(candidate)
        except PermissionError:
            continue

    raise PermissionError(
        f"无法写入结果文件：{target}\n"
        f"已自动尝试追加后缀另存（_1, _2, ...）仍失败。\n"
        f"请检查：\n"
        f"1. 目标文件是否被 Excel/WPS 打开；\n"
        f"2. 输出目录是否受限；\n"
        f"3. 是否可改存到桌面或其他自建文件夹。"
    )


def generate_result(
    source_file,
    source_sheet,
    new_file,
    new_sheet,
    source_primary_col,
    new_primary_col,
    source_secondary_col,
    new_secondary_col,
    match_mode,
    suspect_threshold,
    output_file,
):
    src_headers, src_rows = read_rows(source_file, source_sheet)
    new_headers, new_rows = read_rows(new_file, new_sheet)

    validate_headers(src_headers, "源名单")
    validate_headers(new_headers, "新名单")

    if source_primary_col not in src_headers:
        raise ValueError(f"源名单中未找到主键列：{source_primary_col}")
    if new_primary_col not in new_headers:
        raise ValueError(f"新名单中未找到主键列：{new_primary_col}")

    if source_secondary_col != NO_SECONDARY and source_secondary_col not in src_headers:
        raise ValueError(f"源名单中未找到辅键列：{source_secondary_col}")
    if new_secondary_col != NO_SECONDARY and new_secondary_col not in new_headers:
        raise ValueError(f"新名单中未找到辅键列：{new_secondary_col}")

    if match_mode != MATCH_MODE_PRIMARY:
        if source_secondary_col == NO_SECONDARY or new_secondary_col == NO_SECONDARY:
            raise ValueError("当前匹配模式需要设置辅键列，请先选择源名单和新名单的辅键列。")

    source_metas = build_row_metas(src_rows, source_primary_col, source_secondary_col)
    new_metas = build_row_metas(new_rows, new_primary_col, new_secondary_col)

    source_by_id = {m["id"]: m for m in source_metas}
    new_by_id = {m["id"]: m for m in new_metas}

    source_primary_idx = make_index(source_metas, primary_key)
    new_primary_idx = make_index(new_metas, primary_key)
    source_combined_idx = make_index(source_metas, combined_key)
    new_combined_idx = make_index(new_metas, combined_key)

    warnings = []
    src_primary_empty_count = sum(1 for m in source_metas if not m["primary_norm"])
    new_primary_empty_count = sum(1 for m in new_metas if not m["primary_norm"])
    if src_primary_empty_count:
        warnings.append(f"源名单中有 {src_primary_empty_count} 行主键为空。")
    if new_primary_empty_count:
        warnings.append(f"新名单中有 {new_primary_empty_count} 行主键为空。")

    used_new_ids = set()
    main_rows = []
    unmatched_source_ids = []

    for src in source_metas:
        matched_new_id, status = choose_match_for_source(
            src_meta=src,
            used_new_ids=used_new_ids,
            source_primary_idx=source_primary_idx,
            new_primary_idx=new_primary_idx,
            source_combined_idx=source_combined_idx,
            new_combined_idx=new_combined_idx,
            mode=match_mode,
        )
        matched_new = None
        if matched_new_id is not None:
            used_new_ids.add(matched_new_id)
            matched_new = new_by_id[matched_new_id]
        else:
            unmatched_source_ids.append(src["id"])

        main_rows.append({
            "status": status,
            "src": src,
            "new": matched_new,
        })

    unmatched_new_ids = [m["id"] for m in new_metas if m["id"] not in used_new_ids]
    unmatched_source_metas = [source_by_id[i] for i in unmatched_source_ids]
    unmatched_new_metas = [new_by_id[i] for i in unmatched_new_ids]

    suspects = build_suspect_candidates(
        unmatched_source_metas=unmatched_source_metas,
        unmatched_new_metas=unmatched_new_metas,
        threshold=suspect_threshold,
        topn=3,
    )

    out_wb = Workbook()

    # 汇总
    ws_summary = out_wb.active
    ws_summary.title = "汇总"

    summary = {
        "generated_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "source_file": str(source_file),
        "new_file": str(new_file),
        "match_mode": match_mode,
        "source_primary": source_primary_col,
        "new_primary": new_primary_col,
        "source_secondary": source_secondary_col,
        "new_secondary": new_secondary_col,
        "threshold": suspect_threshold,
        "source_total": len(source_metas),
        "new_total": len(new_metas),
        "count_combined": sum(1 for r in main_rows if r["status"] == STATUS_MATCH_COMBINED),
        "count_primary": sum(1 for r in main_rows if r["status"] == STATUS_MATCH_PRIMARY),
        "count_fallback": sum(1 for r in main_rows if r["status"] == STATUS_MATCH_FALLBACK),
        "count_missing": sum(1 for r in main_rows if r["status"] == STATUS_MISSING),
        "count_primary_dup": sum(1 for r in main_rows if r["status"] == STATUS_PRIMARY_DUP),
        "count_combined_dup": sum(1 for r in main_rows if r["status"] == STATUS_COMBINED_DUP),
        "count_primary_empty": sum(1 for r in main_rows if r["status"] == STATUS_SRC_PRIMARY_EMPTY),
        "count_secondary_empty": sum(1 for r in main_rows if r["status"] == STATUS_SRC_SECONDARY_EMPTY),
        "new_unmatched": len(unmatched_new_ids),
        "suspect_count": len(suspects),
        "warnings": warnings,
    }
    for row in build_summary_rows(summary):
        ws_summary.append(row)
    style_summary_sheet(ws_summary)

    # 主结果
    ws_main = out_wb.create_sheet("按源名单顺序校准")
    main_headers = (
        ["核对状态", "源名单行号", "新名单行号", "源主键(规范化)", "源辅键(规范化)", "新主键(规范化)", "新辅键(规范化)"] +
        [f"源名单_{h}" for h in src_headers] +
        [f"新名单_{h}" for h in new_headers]
    )
    ws_main.append(main_headers)
    for item in main_rows:
        src = item["src"]
        newm = item["new"]
        row = [
            item["status"],
            src["rownum"],
            "" if newm is None else newm["rownum"],
            src["primary_norm"],
            src["secondary_norm"],
            "" if newm is None else newm["primary_norm"],
            "" if newm is None else newm["secondary_norm"],
        ]
        row += [src["data"].get(h) for h in src_headers]
        if newm is None:
            row += [""] * len(new_headers)
        else:
            row += [newm["data"].get(h) for h in new_headers]
        ws_main.append(row)
    style_header(ws_main)
    apply_common_sheet_style(ws_main)
    apply_status_fills(ws_main, status_col=1)
    autofit_columns(ws_main, max_width=35)

    # 源名单缺漏
    ws_missing = out_wb.create_sheet("源名单缺漏")
    missing_headers = ["核对状态", "源名单行号", "源主键(规范化)", "源辅键(规范化)"] + src_headers
    ws_missing.append(missing_headers)
    for item in main_rows:
        if item["status"] in {STATUS_MISSING, STATUS_PRIMARY_DUP, STATUS_COMBINED_DUP, STATUS_SRC_PRIMARY_EMPTY, STATUS_SRC_SECONDARY_EMPTY}:
            src = item["src"]
            ws_missing.append([
                item["status"],
                src["rownum"],
                src["primary_norm"],
                src["secondary_norm"],
            ] + [src["data"].get(h) for h in src_headers])
    style_header(ws_missing)
    apply_common_sheet_style(ws_missing)
    apply_status_fills(ws_missing, status_col=1)
    autofit_columns(ws_missing, max_width=35)

    # 新名单多出
    ws_extra = out_wb.create_sheet("新名单多出")
    extra_headers = ["新名单行号", "新主键(规范化)", "新辅键(规范化)"] + new_headers
    ws_extra.append(extra_headers)
    for nid in unmatched_new_ids:
        nm = new_by_id[nid]
        ws_extra.append([
            nm["rownum"],
            nm["primary_norm"],
            nm["secondary_norm"],
        ] + [nm["data"].get(h) for h in new_headers])
    style_header(ws_extra)
    apply_common_sheet_style(ws_extra)
    autofit_columns(ws_extra, max_width=35)

    # 重复检查
    ws_dup = out_wb.create_sheet("重复检查")
    dup_headers = ["来源", "重复类型", "规范化键", "行号", "主键原值", "辅键原值", "整行内容(JSON)"]
    ws_dup.append(dup_headers)

    for key, ids in sorted(source_primary_idx.items(), key=lambda x: (str(x[0]), len(x[1]))):
        if len(ids) > 1:
            for sid in ids:
                sm = source_by_id[sid]
                ws_dup.append([
                    "源名单",
                    "主键重复",
                    key,
                    sm["rownum"],
                    sm["primary_raw"],
                    sm["secondary_raw"],
                    json.dumps(sm["data"], ensure_ascii=False),
                ])

    for key, ids in sorted(new_primary_idx.items(), key=lambda x: (str(x[0]), len(x[1]))):
        if len(ids) > 1:
            for nid in ids:
                nm = new_by_id[nid]
                ws_dup.append([
                    "新名单",
                    "主键重复",
                    key,
                    nm["rownum"],
                    nm["primary_raw"],
                    nm["secondary_raw"],
                    json.dumps(nm["data"], ensure_ascii=False),
                ])

    for key, ids in sorted(source_combined_idx.items(), key=lambda x: (str(x[0]), len(x[1]))):
        if len(ids) > 1:
            for sid in ids:
                sm = source_by_id[sid]
                ws_dup.append([
                    "源名单",
                    "联合键重复",
                    f"{key[0]} | {key[1]}",
                    sm["rownum"],
                    sm["primary_raw"],
                    sm["secondary_raw"],
                    json.dumps(sm["data"], ensure_ascii=False),
                ])

    for key, ids in sorted(new_combined_idx.items(), key=lambda x: (str(x[0]), len(x[1]))):
        if len(ids) > 1:
            for nid in ids:
                nm = new_by_id[nid]
                ws_dup.append([
                    "新名单",
                    "联合键重复",
                    f"{key[0]} | {key[1]}",
                    nm["rownum"],
                    nm["primary_raw"],
                    nm["secondary_raw"],
                    json.dumps(nm["data"], ensure_ascii=False),
                ])
    style_header(ws_dup)
    apply_common_sheet_style(ws_dup)
    autofit_columns(ws_dup, max_width=60)

    # 疑似匹配
    ws_suspect = out_wb.create_sheet("疑似匹配建议")
    suspect_headers = [
        "建议序号",
        "相似度得分",
        "源名单行号",
        "新名单行号",
        "源主键原值",
        "新主键原值",
        "源主键(规范化)",
        "新主键(规范化)",
        "源辅键原值",
        "新辅键原值",
        "源辅键(规范化)",
        "新辅键(规范化)",
    ]
    ws_suspect.append(suspect_headers)
    for item in suspects:
        sm = item["source_meta"]
        nm = item["new_meta"]
        ws_suspect.append([
            item["rank"],
            item["score"],
            sm["rownum"],
            nm["rownum"],
            sm["primary_raw"],
            nm["primary_raw"],
            sm["primary_norm"],
            nm["primary_norm"],
            sm["secondary_raw"],
            nm["secondary_raw"],
            sm["secondary_norm"],
            nm["secondary_norm"],
        ])
    style_header(ws_suspect)
    apply_common_sheet_style(ws_suspect)
    autofit_columns(ws_suspect, max_width=35)

    # 说明
    ws_note = out_wb.create_sheet("说明")
    notes = [
        ["本工具完全离线运行，不上传任何数据。"],
        ["v3 满血版功能：自动改名保存、主键+辅键联合匹配、疑似匹配建议、结果高亮、统计汇总、重复检查。"],
        ["主键一般选择：姓名。"],
        ["辅键建议选择：电话后4位、身份证后4位、年龄、村/社区等。"],
        ["匹配模式说明："],
        [f"1. {MATCH_MODE_PRIMARY}：只按主键匹配；若主键在任一表中重复，则不自动匹配。"],
        [f"2. {MATCH_MODE_COMBINED_STRICT}：要求主键和辅键都一致；更稳，但辅键为空时不会自动匹配。"],
        [f"3. {MATCH_MODE_COMBINED_FALLBACK}：优先用主键+辅键匹配；匹配不到时，如果主键在两表中都唯一，则回退按主键匹配。"],
        ["疑似匹配建议：仅用于人工复核，不会自动合并。"],
        ["若结果文件正被 Excel/WPS 占用，程序会自动尝试另存为 _1、_2、_3 ...。"],
    ]
    for row in notes:
        ws_note.append(row)
    ws_note.column_dimensions["A"].width = 120

    saved_path = safe_save_workbook(out_wb, output_file)
    return saved_path


class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title(APP_TITLE)
        self.geometry("980x680")
        self.resizable(True, True)

        self.source_file_var = tk.StringVar()
        self.new_file_var = tk.StringVar()
        self.output_file_var = tk.StringVar()

        self.source_sheet_var = tk.StringVar()
        self.new_sheet_var = tk.StringVar()

        self.source_primary_var = tk.StringVar(value="姓名")
        self.new_primary_var = tk.StringVar(value="姓名")
        self.source_secondary_var = tk.StringVar(value=NO_SECONDARY)
        self.new_secondary_var = tk.StringVar(value=NO_SECONDARY)

        self.match_mode_var = tk.StringVar(value=MATCH_MODE_COMBINED_FALLBACK)
        self.threshold_var = tk.StringVar(value="0.66")

        self._build_ui()

    def _build_ui(self):
        pad = {"padx": 8, "pady": 6}

        title = ttk.Label(self, text=APP_TITLE, font=("Microsoft YaHei UI", 16, "bold"))
        title.grid(row=0, column=0, columnspan=6, sticky="w", **pad)

        ttk.Label(self, text="1）源名单 Excel：").grid(row=1, column=0, sticky="e", **pad)
        ttk.Entry(self, textvariable=self.source_file_var, width=78).grid(row=1, column=1, columnspan=3, sticky="we", **pad)
        ttk.Button(self, text="浏览", command=self.pick_source).grid(row=1, column=4, sticky="w", **pad)

        ttk.Label(self, text="源名单工作表：").grid(row=2, column=0, sticky="e", **pad)
        self.source_sheet_cb = ttk.Combobox(self, textvariable=self.source_sheet_var, state="readonly", width=26)
        self.source_sheet_cb.grid(row=2, column=1, sticky="w", **pad)
        self.source_sheet_cb.bind("<<ComboboxSelected>>", lambda e: self.load_source_headers())

        ttk.Label(self, text="源主键列：").grid(row=2, column=2, sticky="e", **pad)
        self.source_primary_cb = ttk.Combobox(self, textvariable=self.source_primary_var, state="readonly", width=22)
        self.source_primary_cb.grid(row=2, column=3, sticky="w", **pad)

        ttk.Label(self, text="源辅键列：").grid(row=2, column=4, sticky="e", **pad)
        self.source_secondary_cb = ttk.Combobox(self, textvariable=self.source_secondary_var, state="readonly", width=22)
        self.source_secondary_cb.grid(row=2, column=5, sticky="w", **pad)

        ttk.Label(self, text="2）新名单 Excel：").grid(row=3, column=0, sticky="e", **pad)
        ttk.Entry(self, textvariable=self.new_file_var, width=78).grid(row=3, column=1, columnspan=3, sticky="we", **pad)
        ttk.Button(self, text="浏览", command=self.pick_new).grid(row=3, column=4, sticky="w", **pad)

        ttk.Label(self, text="新名单工作表：").grid(row=4, column=0, sticky="e", **pad)
        self.new_sheet_cb = ttk.Combobox(self, textvariable=self.new_sheet_var, state="readonly", width=26)
        self.new_sheet_cb.grid(row=4, column=1, sticky="w", **pad)
        self.new_sheet_cb.bind("<<ComboboxSelected>>", lambda e: self.load_new_headers())

        ttk.Label(self, text="新主键列：").grid(row=4, column=2, sticky="e", **pad)
        self.new_primary_cb = ttk.Combobox(self, textvariable=self.new_primary_var, state="readonly", width=22)
        self.new_primary_cb.grid(row=4, column=3, sticky="w", **pad)

        ttk.Label(self, text="新辅键列：").grid(row=4, column=4, sticky="e", **pad)
        self.new_secondary_cb = ttk.Combobox(self, textvariable=self.new_secondary_var, state="readonly", width=22)
        self.new_secondary_cb.grid(row=4, column=5, sticky="w", **pad)

        ttk.Label(self, text="3）匹配模式：").grid(row=5, column=0, sticky="e", **pad)
        self.mode_cb = ttk.Combobox(
            self,
            textvariable=self.match_mode_var,
            state="readonly",
            values=[MATCH_MODE_PRIMARY, MATCH_MODE_COMBINED_STRICT, MATCH_MODE_COMBINED_FALLBACK],
            width=32,
        )
        self.mode_cb.grid(row=5, column=1, sticky="w", **pad)

        ttk.Label(self, text="疑似匹配阈值：").grid(row=5, column=2, sticky="e", **pad)
        ttk.Entry(self, textvariable=self.threshold_var, width=10).grid(row=5, column=3, sticky="w", **pad)
        ttk.Label(self, text="建议 0.60 ~ 0.80，默认 0.66").grid(row=5, column=4, columnspan=2, sticky="w", **pad)

        ttk.Label(self, text="4）导出结果文件：").grid(row=6, column=0, sticky="e", **pad)
        ttk.Entry(self, textvariable=self.output_file_var, width=78).grid(row=6, column=1, columnspan=3, sticky="we", **pad)
        ttk.Button(self, text="选择保存位置", command=self.pick_output).grid(row=6, column=4, sticky="w", **pad)

        tips = (
            "使用建议：\n"
            "• 普通场景：主键选“姓名”，匹配模式选“优先联合匹配，必要时回退主键”。\n"
            "• 重名较多场景：辅键选“电话后4位 / 身份证后4位 / 年龄 / 村社区”等，再选“主键+辅键联合精确匹配”。\n"
            "• 输出包含：汇总、按源名单顺序校准、源名单缺漏、新名单多出、重复检查、疑似匹配建议、说明。\n"
            "• 若目标文件被占用，程序会自动尝试保存为 _1、_2、_3 ...。\n"
            "• 支持 .xlsx / .xlsm；旧版 .xls 请先另存为 .xlsx。"
        )
        self.tip_text = tk.Text(self, height=13, width=110)
        self.tip_text.grid(row=7, column=0, columnspan=6, sticky="nsew", padx=10, pady=10)
        self.tip_text.insert("1.0", tips)
        self.tip_text.config(state="disabled")

        self.run_btn = ttk.Button(self, text="开始校准并导出", command=self.run_align)
        self.run_btn.grid(row=8, column=0, columnspan=6, pady=12)

        self.status_var = tk.StringVar(value="就绪")
        ttk.Label(self, textvariable=self.status_var, foreground="blue").grid(row=9, column=0, columnspan=6, sticky="w", padx=10, pady=8)

        self.columnconfigure(1, weight=1)
        self.columnconfigure(3, weight=1)
        self.rowconfigure(7, weight=1)

    def pick_source(self):
        path = filedialog.askopenfilename(
            title="选择源名单 Excel",
            filetypes=[("Excel 文件", "*.xlsx *.xlsm"), ("所有文件", "*.*")]
        )
        if not path:
            return
        self.source_file_var.set(path)
        names = get_sheet_names(path)
        self.source_sheet_cb["values"] = names
        if names:
            self.source_sheet_var.set(names[0])
            self.load_source_headers()
        if not self.output_file_var.get():
            self.output_file_var.set(str(Path(path).with_name("名单校准结果_v3.xlsx")))

    def pick_new(self):
        path = filedialog.askopenfilename(
            title="选择新名单 Excel",
            filetypes=[("Excel 文件", "*.xlsx *.xlsm"), ("所有文件", "*.*")]
        )
        if not path:
            return
        self.new_file_var.set(path)
        names = get_sheet_names(path)
        self.new_sheet_cb["values"] = names
        if names:
            self.new_sheet_var.set(names[0])
            self.load_new_headers()
        if not self.output_file_var.get():
            self.output_file_var.set(str(Path(path).with_name("名单校准结果_v3.xlsx")))

    def pick_output(self):
        path = filedialog.asksaveasfilename(
            title="保存结果为",
            defaultextension=".xlsx",
            initialfile="名单校准结果_v3.xlsx",
            filetypes=[("Excel 文件", "*.xlsx")]
        )
        if path:
            self.output_file_var.set(path)

    def load_source_headers(self):
        fp = self.source_file_var.get().strip()
        sh = self.source_sheet_var.get().strip()
        if not fp or not sh:
            return
        headers = read_sheet_headers(fp, sh)
        self.source_primary_cb["values"] = headers
        self.source_secondary_cb["values"] = [NO_SECONDARY] + headers
        if "姓名" in headers:
            self.source_primary_var.set("姓名")
        elif headers:
            self.source_primary_var.set(headers[0])
        self.source_secondary_var.set(NO_SECONDARY)

    def load_new_headers(self):
        fp = self.new_file_var.get().strip()
        sh = self.new_sheet_var.get().strip()
        if not fp or not sh:
            return
        headers = read_sheet_headers(fp, sh)
        self.new_primary_cb["values"] = headers
        self.new_secondary_cb["values"] = [NO_SECONDARY] + headers
        if "姓名" in headers:
            self.new_primary_var.set("姓名")
        elif headers:
            self.new_primary_var.set(headers[0])
        self.new_secondary_var.set(NO_SECONDARY)

    def run_align(self):
        source_file = self.source_file_var.get().strip()
        new_file = self.new_file_var.get().strip()
        source_sheet = self.source_sheet_var.get().strip()
        new_sheet = self.new_sheet_var.get().strip()
        output_file = self.output_file_var.get().strip()

        if not source_file or not Path(source_file).exists():
            messagebox.showerror("错误", "请先选择有效的源名单 Excel 文件。")
            return
        if not new_file or not Path(new_file).exists():
            messagebox.showerror("错误", "请先选择有效的新名单 Excel 文件。")
            return
        if not source_sheet:
            messagebox.showerror("错误", "请选择源名单工作表。")
            return
        if not new_sheet:
            messagebox.showerror("错误", "请选择新名单工作表。")
            return
        if not output_file:
            messagebox.showerror("错误", "请选择结果文件保存位置。")
            return

        try:
            threshold = float(self.threshold_var.get().strip())
        except Exception:
            messagebox.showerror("错误", "疑似匹配阈值必须是数字，例如 0.66。")
            return

        if not (0 <= threshold <= 1):
            messagebox.showerror("错误", "疑似匹配阈值必须在 0 到 1 之间。")
            return

        try:
            self.status_var.set("处理中，请稍候……")
            self.update_idletasks()

            saved_path = generate_result(
                source_file=source_file,
                source_sheet=source_sheet,
                new_file=new_file,
                new_sheet=new_sheet,
                source_primary_col=self.source_primary_var.get().strip(),
                new_primary_col=self.new_primary_var.get().strip(),
                source_secondary_col=self.source_secondary_var.get().strip() or NO_SECONDARY,
                new_secondary_col=self.new_secondary_var.get().strip() or NO_SECONDARY,
                match_mode=self.match_mode_var.get().strip(),
                suspect_threshold=threshold,
                output_file=output_file,
            )
            self.status_var.set(f"完成：{saved_path}")
            messagebox.showinfo("完成", f"已成功导出结果：\n{saved_path}")
        except Exception as e:
            self.status_var.set("处理失败")
            detail = traceback.format_exc()
            messagebox.showerror("处理失败", f"{e}\n\n详细信息：\n{detail}")


if __name__ == "__main__":
    app = App()
    app.mainloop()
